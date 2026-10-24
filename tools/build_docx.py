"""Build the editable report from docs/report.md and application wireframes."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]

def set_font(run,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(14);run.font.color.rgb=RGBColor(0,0,0);run.bold=bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')

def rich(p,text,bold=False):
    text=re.sub(r'`([^`]+)`',r'\1',text)
    for part in re.split(r'(https?://[^\s]+)',text):
        if part.startswith(('http://','https://')):
            url=part.rstrip('.,;');tail=part[len(url):]
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
            for name,attrs in [('rFonts',{'ascii':'Times New Roman','hAnsi':'Times New Roman'}),('sz',{'val':'28'}),('color',{'val':'0000FF'}),('u',{'val':'single'})]:
                e=OxmlElement('w:'+name)
                for k,v in attrs.items():e.set(qn('w:'+k),v)
                pr.append(e)
            r.append(pr);t=OxmlElement('w:t');t.text=url;r.append(t);link.append(r);p._p.append(link)
            if tail:set_font(p.add_run(tail),bold)
        elif part:
            for k,seg in enumerate(re.split(r'\*\*(.*?)\*\*',part)):
                if seg:set_font(p.add_run(seg),bold or k%2==1)

def field(p,instruction):
    r=p.add_run();set_font(r)
    begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=instruction
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    for e in [begin,instr,sep,end]:r._r.append(e)

def build():
    doc=Document();doc.core_properties.author='Yan Kudashov';doc.core_properties.title='ASTI First Probationary Period Report';sec=doc.sections[0];sec.page_width=Pt(595.276);sec.page_height=Pt(841.89)
    sec.left_margin=sec.right_margin=Pt(54);sec.top_margin=Pt(55);sec.bottom_margin=Pt(49)
    sec.header_distance=Pt(15);sec.footer_distance=Pt(20)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','TOC 1','TOC Heading','Hyperlink']:
        if name not in doc.styles:
            from docx.enum.style import WD_STYLE_TYPE
            doc.styles.add_style(name,WD_STYLE_TYPE.CHARACTER if name=='Hyperlink' else WD_STYLE_TYPE.PARAGRAPH)
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(14);st.font.color.rgb=RGBColor(0,0,255) if name=='Hyperlink' else RGBColor(0,0,0)
        if name!='Hyperlink':
            st.paragraph_format.line_spacing=1.5;st.paragraph_format.space_after=Pt(8)
            st.paragraph_format.widow_control=True
        if name.startswith('Heading'):
            st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12)
    doc.styles['Title'].font.bold=True
    for st in doc.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    settings=doc.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;field(foot,'PAGE')
    doc.add_paragraph('ASTI',style='Title')
    doc.add_paragraph('First Probationary Period Report',style='Subtitle')
    doc.add_paragraph('Development of an Intelligent System for a Physical Desktop Assistant with a Mobile Application and External Service Integration',style='Subtitle')
    for t in ['Yan Kudashov | SE-2413','Supervisor: Сарсенова Жибек Нуржановна','Astana IT University','Software Engineering','24 October 2026']:
        rich(doc.add_paragraph(),t)
    for t in ['Repository: https://github.com/Kreichik/asti','YouTrack: https://asti-diploma.youtrack.cloud/agiles/204-1/current']:rich(doc.add_paragraph(),t)
    doc.add_page_break();doc.add_paragraph('Contents',style='TOC Heading');field(doc.add_paragraph(),'TOC \\o "1-1" \\h \\z \\u');doc.add_page_break()
    lines=(ROOT/'docs/report.md').read_text(encoding='utf-8').splitlines();i=next(i for i,l in enumerate(lines) if l.startswith('## 1 '))
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('## '):doc.add_heading(line[3:],level=1);i+=1;continue
        if line.startswith('### '):doc.add_heading(line[4:],level=2);i+=1;continue
        if line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(ROOT/'docs'/m[2]),width=Pt(487))
            doc.add_paragraph(m[1],style='Caption');i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.match(r'^:?-+:?$',c) for c in cells):rows.append(cells)
                i+=1
            n=len(rows[0]);widths={3:[62,65,360],4:[92,132,132,131],5:[84,101,101,101,100]}.get(n,[487/n]*n)
            if n==3 and rows[0][0] in ['Participants','Area','Milestone','Risk']:widths=[95,196,196]
            if n==3 and rows[0][0]=='Alternative':widths=[100,200,187]
            table=doc.add_table(rows=0,cols=n);table.alignment=WD_TABLE_ALIGNMENT.LEFT;table.autofit=False
            pr=table._tbl.tblPr;borders=OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                e=OxmlElement('w:'+edge)
                for k,v in [('val','single'),('sz','4'),('color','000000')]:e.set(qn('w:'+k),v)
                borders.append(e)
            pr.append(borders)
            for col,width in zip(table.columns,widths):col.width=Pt(width)
            for ri,values in enumerate(rows):
                row=table.add_row();trpr=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');trpr.append(cs)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
                for j,val in enumerate(values):
                    cell=row.cells[j];cell.width=Pt(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cp=cell._tc.get_or_add_tcPr()
                    cell_borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        edge_element=OxmlElement('w:'+edge)
                        for key,value in [('val','single'),('sz','4'),('color','000000')]:edge_element.set(qn('w:'+key),value)
                        cell_borders.append(edge_element)
                    cp.append(cell_borders)
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FFFFFF');cp.append(shade)
                    marg=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');marg.append(e)
                    cp.append(marg)
                    p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(0);rich(p,val,ri==0)
            doc.add_paragraph();continue
        para=[line];i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','![')):
            para.append(lines[i].strip());i+=1
        rich(doc.add_paragraph(),' '.join(para))
    out=ROOT/'output/docx/SE-2413_Kudashov_ProbPeriod1.docx';out.parent.mkdir(parents=True,exist_ok=True);doc.save(out);print(out)
if __name__=='__main__':build()
